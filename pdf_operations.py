"""Operaciones de PDF independientes de la interfaz gráfica."""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from difflib import unified_diff

import fitz
from pdf2docx import Converter
from PIL import Image
import pytesseract
import xlsxwriter
from pptx import Presentation
from pptx.util import Inches
from PyPDF2 import PdfMerger, PdfReader, PdfWriter


def _require_pdf(source: str | Path) -> Path:
    path = Path(source)
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo PDF: {path}")
    return path


def _prepare_destination(destination: str | Path, *sources: str | Path) -> Path:
    path = Path(destination)
    source_paths = {Path(source).resolve() for source in sources}
    if path.resolve() in source_paths:
        raise ValueError("El archivo de salida debe ser distinto de los archivos de entrada.")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def convert_pdf_to_word(source: str | Path, destination: str | Path) -> None:
    """Convierte un PDF a DOCX y cierra siempre el conversor."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    converter = Converter(str(source_path))
    try:
        converter.convert(str(destination_path), start=0, end=None)
    finally:
        converter.close()


def merge_pdfs(sources: list[str | Path], destination: str | Path) -> None:
    """Une dos o más PDFs en el orden recibido."""
    if len(sources) < 2:
        raise ValueError("Se necesitan al menos dos archivos PDF.")

    source_paths = [_require_pdf(source) for source in sources]
    destination_path = _prepare_destination(destination, *source_paths)
    merger = PdfMerger()
    try:
        for source in source_paths:
            merger.append(str(source))
        merger.write(str(destination_path))
    finally:
        merger.close()


def split_pdf(source: str | Path, output_directory: str | Path) -> list[Path]:
    """Extrae cada página a un archivo page_N.pdf."""
    source_path = _require_pdf(source)
    output_dir = Path(output_directory)
    output_dir.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(source_path))
    outputs: list[Path] = []

    for page_number, page in enumerate(reader.pages, start=1):
        destination = output_dir / f"page_{page_number}.pdf"
        writer = PdfWriter()
        writer.add_page(page)
        with destination.open("wb") as output_file:
            writer.write(output_file)
        outputs.append(destination)

    return outputs


def convert_pdf_to_powerpoint(source: str | Path, destination: str | Path) -> None:
    """Crea una presentación con una imagen por página del PDF."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)

    with TemporaryDirectory(prefix="pdftoolkit-") as temporary_directory:
        temporary_path = Path(temporary_directory)
        document = fitz.open(str(source_path))
        try:
            for page_number, page in enumerate(document, start=1):
                image_path = temporary_path / f"page_{page_number}.png"
                page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(str(image_path))

                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                page_ratio = page.rect.width / page.rect.height
                slide_ratio = 10 / 7.5
                if page_ratio >= slide_ratio:
                    width = 10
                    height = width / page_ratio
                else:
                    height = 7.5
                    width = height * page_ratio
                slide.shapes.add_picture(
                    str(image_path),
                    Inches((10 - width) / 2),
                    Inches((7.5 - height) / 2),
                    width=Inches(width),
                    height=Inches(height),
                )
        finally:
            document.close()

    presentation.save(str(destination_path))


def compress_pdf(source: str | Path, destination: str | Path) -> None:
    """Guarda una copia optimizada del PDF usando la compresión de MuPDF."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    document = fitz.open(str(source_path))
    try:
        document.save(str(destination_path), garbage=4, clean=True, deflate=True)
    finally:
        document.close()


def rotate_pdf(source: str | Path, destination: str | Path, angle: int) -> None:
    """Rota todas las páginas 90, 180 o 270 grados."""
    if angle not in (90, 180, 270):
        raise ValueError("El ángulo debe ser 90, 180 o 270 grados.")
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    reader = PdfReader(str(source_path))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page.rotate(angle))
    with destination_path.open("wb") as output_file:
        writer.write(output_file)


def pdf_to_jpg(source: str | Path, output_directory: str | Path, dpi: int = 150) -> list[Path]:
    """Renderiza cada página del PDF como una imagen JPG."""
    if dpi < 72 or dpi > 600:
        raise ValueError("El DPI debe estar entre 72 y 600.")
    source_path = _require_pdf(source)
    output_dir = Path(output_directory)
    output_dir.mkdir(parents=True, exist_ok=True)
    scale = dpi / 72
    outputs: list[Path] = []
    document = fitz.open(str(source_path))
    try:
        for page_number, page in enumerate(document, start=1):
            destination = output_dir / f"page_{page_number}.jpg"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            pixmap.save(str(destination), output="jpeg")
            outputs.append(destination)
    finally:
        document.close()
    return outputs


def jpg_to_pdf(sources: list[str | Path], destination: str | Path) -> None:
    """Convierte una o más imágenes JPG/PNG en un PDF."""
    if not sources:
        raise ValueError("Selecciona al menos una imagen.")
    source_paths = [Path(source) for source in sources]
    for source_path in source_paths:
        if not source_path.is_file():
            raise FileNotFoundError(f"No existe la imagen: {source_path}")
    destination_path = _prepare_destination(destination, *source_paths)
    images: list[Image.Image] = []
    try:
        for source_path in source_paths:
            with Image.open(source_path) as image:
                images.append(image.convert("RGB"))
        images[0].save(destination_path, "PDF", save_all=True, append_images=images[1:])
    finally:
        for image in images:
            image.close()


def add_watermark(source: str | Path, destination: str | Path, text: str) -> None:
    """Añade una marca de agua sencilla a todas las páginas."""
    if not text.strip():
        raise ValueError("El texto de la marca de agua no puede estar vacío.")
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    document = fitz.open(str(source_path))
    try:
        for page in document:
            center = page.rect.width / 2, page.rect.height / 2
            page.insert_text(
                center,
                text,
                fontsize=min(page.rect.width, page.rect.height) / 10,
                color=(0.55, 0.55, 0.55),
                overlay=True,
            )
        document.save(str(destination_path), garbage=4, deflate=True)
    finally:
        document.close()


def protect_pdf(source: str | Path, destination: str | Path, password: str) -> None:
    """Protege un PDF con contraseña de usuario."""
    if not password:
        raise ValueError("La contraseña no puede estar vacía.")
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    reader = PdfReader(str(source_path))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(password)
    with destination_path.open("wb") as output_file:
        writer.write(output_file)


def unlock_pdf(source: str | Path, destination: str | Path, password: str) -> None:
    """Quita la protección usando la contraseña válida del documento."""
    if not password:
        raise ValueError("Introduce la contraseña del PDF protegido.")
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    reader = PdfReader(str(source_path))
    if reader.is_encrypted and reader.decrypt(password) == 0:
        raise ValueError("La contraseña del PDF no es válida.")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    with destination_path.open("wb") as output_file:
        writer.write(output_file)


def reorder_pdf(source: str | Path, destination: str | Path, order: list[int]) -> None:
    """Reordena páginas usando una lista numerada desde 1."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    reader = PdfReader(str(source_path))
    expected = list(range(1, len(reader.pages) + 1))
    if sorted(order) != expected:
        raise ValueError(f"El orden debe contener exactamente las páginas: {expected}.")
    writer = PdfWriter()
    for page_number in order:
        writer.add_page(reader.pages[page_number - 1])
    with destination_path.open("wb") as output_file:
        writer.write(output_file)


def number_pdf(source: str | Path, destination: str | Path) -> None:
    """Añade el número de página en el pie de cada página."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    document = fitz.open(str(source_path))
    try:
        for page_number, page in enumerate(document, start=1):
            page.insert_text(
                (page.rect.width - 55, page.rect.height - 24),
                str(page_number),
                fontsize=10,
                color=(0.25, 0.25, 0.25),
                overlay=True,
            )
        document.save(str(destination_path), garbage=4, deflate=True)
    finally:
        document.close()


def ocr_pdf(source: str | Path, destination: str | Path, language: str = "eng") -> None:
    """Extrae texto de cada página renderizada mediante Tesseract OCR."""
    if not language.strip():
        raise ValueError("Indica al menos un idioma OCR, por ejemplo: eng o spa+eng.")
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    try:
        pytesseract.get_tesseract_version()
    except Exception as error:
        raise RuntimeError(
            "Tesseract OCR no está instalado o no está disponible en el PATH de Windows."
        ) from error

    document = fitz.open(str(source_path))
    try:
        extracted_pages: list[str] = []
        matrix = fitz.Matrix(2, 2)
        for page_number, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            text = pytesseract.image_to_string(image, lang=language)
            extracted_pages.append(f"--- Página {page_number} ---\n{text.strip()}\n")
            image.close()
    finally:
        document.close()

    destination_path.write_text("\n".join(extracted_pages), encoding="utf-8")


def pdf_to_excel(source: str | Path, destination: str | Path) -> None:
    """Exporta el texto de cada página a una hoja independiente de Excel."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    workbook = xlsxwriter.Workbook(str(destination_path))
    try:
        header_format = workbook.add_format({"bold": True, "bg_color": "#DCE6F1"})
        document = fitz.open(str(source_path))
        try:
            for page_number, page in enumerate(document, start=1):
                worksheet = workbook.add_worksheet(f"Página {page_number}"[:31])
                worksheet.write(0, 0, "Línea", header_format)
                worksheet.write(0, 1, "Texto", header_format)
                lines = page.get_text("text").splitlines()
                for line_number, line in enumerate(lines, start=1):
                    worksheet.write(line_number, 0, line_number)
                    worksheet.write(line_number, 1, line)
                worksheet.set_column("A:A", 10)
                worksheet.set_column("B:B", 100)
        finally:
            document.close()
    finally:
        workbook.close()


def edit_pdf_metadata(
    source: str | Path,
    destination: str | Path,
    title: str = "",
    author: str = "",
    subject: str = "",
) -> None:
    """Edita metadatos básicos sin alterar el contenido de las páginas."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    document = fitz.open(str(source_path))
    try:
        metadata = document.metadata
        metadata.update({"title": title, "author": author, "subject": subject})
        document.set_metadata(metadata)
        document.save(str(destination_path), garbage=4, deflate=True)
    finally:
        document.close()


def repair_pdf(source: str | Path, destination: str | Path) -> None:
    """Reconstruye y guarda el PDF para eliminar objetos dañados o redundantes."""
    source_path = _require_pdf(source)
    destination_path = _prepare_destination(destination, source_path)
    document = fitz.open(str(source_path))
    try:
        document.save(str(destination_path), garbage=4, clean=True, deflate=True)
    finally:
        document.close()


def compare_pdfs(first: str | Path, second: str | Path, destination: str | Path) -> None:
    """Genera un informe de diferencias de texto entre dos PDFs."""
    first_path = _require_pdf(first)
    second_path = _require_pdf(second)
    destination_path = _prepare_destination(destination, first_path, second_path)
    first_document = fitz.open(str(first_path))
    second_document = fitz.open(str(second_path))
    try:
        first_lines = "\n".join(page.get_text("text") for page in first_document).splitlines()
        second_lines = "\n".join(page.get_text("text") for page in second_document).splitlines()
    finally:
        first_document.close()
        second_document.close()

    differences = unified_diff(
        first_lines,
        second_lines,
        fromfile=str(first_path.name),
        tofile=str(second_path.name),
        lineterm="",
    )
    destination_path.write_text("\n".join(differences), encoding="utf-8")
